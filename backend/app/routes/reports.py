import csv
import io

from flask import Blueprint, Response, request

from ..db import db_cursor
from ..responses import fail, ok
from ..security import ROLE_ADMIN, require_roles

reports_bp = Blueprint("reports", __name__)


def filters():
    where = []
    params = []
    start = request.args.get("fecha_inicio")
    end = request.args.get("fecha_fin")
    barber_id = request.args.get("barbero_id", type=int)
    status = request.args.get("estado")
    service_id = request.args.get("servicio_id", type=int)
    if start:
        where.append("r.fecha >= %s")
        params.append(start)
    if end:
        where.append("r.fecha <= %s")
        params.append(end)
    if barber_id:
        where.append("r.id_barbero = %s")
        params.append(barber_id)
    if status:
        where.append("r.estado = %s")
        params.append(status)
    if service_id:
        where.append("rs.id_servicio = %s")
        params.append(service_id)
    return ("WHERE " + " AND ".join(where)) if where else "", params


def run_query(cursor, query, params):
    if params:
        cursor.execute(query.replace("%Y", "%%Y").replace("%m", "%%m"), params)
    else:
        cursor.execute(query)


def collect_summary(cursor, where, params):
    no_show_condition = "LEFT(estado, 9) = 'No asisti'"

    run_query(
        cursor,
        f"""
        SELECT
          COUNT(*) AS total,
          SUM(CASE WHEN base.estado = 'Confirmada' THEN 1 ELSE 0 END) AS confirmadas,
          SUM(CASE WHEN base.estado = 'Atendida' THEN 1 ELSE 0 END) AS atendidas,
          SUM(CASE WHEN base.estado = 'Cancelada' THEN 1 ELSE 0 END) AS canceladas,
          SUM(CASE WHEN {no_show_condition.replace('estado', 'base.estado')} THEN 1 ELSE 0 END) AS no_show,
          SUM(CASE WHEN base.estado = 'Atendida' THEN base.total ELSE 0 END) AS ingresos
        FROM (
          SELECT DISTINCT r.id, r.estado, r.total
          FROM reservas r
          LEFT JOIN reserva_servicios rs ON rs.id_reserva = r.id
          {where}
        ) base
        """,
        params,
    )
    summary = cursor.fetchone()

    run_query(
        cursor,
        f"""
        SELECT DATE_FORMAT(base.fecha, '%Y-%m') AS mes,
               SUM(CASE WHEN base.estado = 'Atendida' THEN 1 ELSE 0 END) AS atendidas,
               SUM(CASE WHEN {no_show_condition.replace('estado', 'base.estado')} THEN 1 ELSE 0 END) AS no_show
        FROM (
          SELECT DISTINCT r.id, r.fecha, r.estado
          FROM reservas r
          LEFT JOIN reserva_servicios rs ON rs.id_reserva = r.id
          {where}
        ) base
        GROUP BY DATE_FORMAT(base.fecha, '%Y-%m')
        ORDER BY mes
        """,
        params,
    )
    monthly = cursor.fetchall()

    run_query(
        cursor,
        f"""
        SELECT base.barbero, COUNT(*) AS total,
               SUM(CASE WHEN base.estado = 'Atendida' THEN base.total ELSE 0 END) AS ingresos
        FROM (
          SELECT DISTINCT r.id, b.nombre AS barbero, r.estado, r.total
          FROM reservas r
          INNER JOIN barberos b ON b.id = r.id_barbero
          LEFT JOIN reserva_servicios rs ON rs.id_reserva = r.id
          {where}
        ) base
        GROUP BY base.barbero
        ORDER BY total DESC
        """,
        params,
    )
    by_barber = cursor.fetchall()

    run_query(
        cursor,
        f"""
        SELECT s.nombre AS servicio, COUNT(*) AS total
        FROM reservas r
        INNER JOIN reserva_servicios rs ON rs.id_reserva = r.id
        INNER JOIN servicios s ON s.id = rs.id_servicio
        {where}
        GROUP BY s.nombre
        ORDER BY total DESC
        LIMIT 8
        """,
        params,
    )
    by_service = cursor.fetchall()

    cursor.execute(
        """
        SELECT MAX(CONCAT_WS(' ', c.nombre, NULLIF(c.apellido, ''))) AS nombre, c.telefono,
               COUNT(r.id) AS total_citas,
               SUM(CASE WHEN LEFT(r.estado, 9) = 'No asisti' THEN 1 ELSE 0 END) AS strikes
        FROM clientes c
        LEFT JOIN reservas r ON r.id_cliente = c.id
        GROUP BY c.telefono
        HAVING total_citas > 0
        ORDER BY total_citas DESC
        LIMIT 10
        """
    )
    frequent_clients = cursor.fetchall()

    cursor.execute(
        """
        SELECT id, nombre, cantidad, stock_minimo
        FROM inventario
        WHERE estado = 'Activo' AND cantidad <= stock_minimo
        ORDER BY cantidad ASC
        """
    )
    low_stock = cursor.fetchall()

    cursor.execute("SELECT SUM(total) AS ingresos_productos FROM ventas_inventario")
    inventory_income = cursor.fetchone()

    return {
        "totalReservas": int(summary["total"] or 0),
        "reservasConfirmadas": int(summary["confirmadas"] or 0),
        "reservasCompletadas": int(summary["atendidas"] or 0),
        "reservasCanceladas": int(summary["canceladas"] or 0),
        "reservasNoShow": int(summary["no_show"] or 0),
        "ingresosTotales": float(summary["ingresos"] or 0),
        "ingresosProductos": float(inventory_income["ingresos_productos"] or 0),
        "mensual": monthly,
        "porBarbero": by_barber,
        "porServicio": by_service,
        "clientesFrecuentes": frequent_clients,
        "stockBajo": low_stock,
    }


def summary_payload(cursor, where, params):
    return collect_summary(cursor, where, params)


@reports_bp.get("/summary")
@require_roles(ROLE_ADMIN)
def report_summary():
    """Calcula estadisticas, ingresos, inventario y datos para graficas."""
    where, params = filters()
    with db_cursor() as cursor:
        return ok(collect_summary(cursor, where, params))


def rows_as_lists(rows, columns):
    return [[row.get(key) for key in columns] for row in rows]


def simple_pdf_response(lines):
    """Genera un PDF basico sin dependencias externas como respaldo."""
    escaped_lines = []
    for line in lines:
        safe = str(line).replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        escaped_lines.append(safe[:110])

    stream_parts = ["BT", "/F1 10 Tf", "50 760 Td"]
    for index, line in enumerate(escaped_lines):
        if index:
            stream_parts.append("0 -14 Td")
        stream_parts.append(f"({line}) Tj")
    stream_parts.append("ET")
    stream = "\n".join(stream_parts).encode("latin-1", "replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = io.BytesIO()
    output.write(b"%PDF-1.4\n")
    offsets = [0]
    for number, content in enumerate(objects, start=1):
        offsets.append(output.tell())
        output.write(f"{number} 0 obj\n".encode("ascii"))
        output.write(content)
        output.write(b"\nendobj\n")
    xref = output.tell()
    output.write(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.write(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.write(f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode("ascii"))
    return Response(
        output.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": "attachment; filename=reporte-barberia.pdf"},
    )


def fetch_reservation_rows(cursor, where, params):
    cursor.execute(
        f"""
        SELECT r.fecha, TIME_FORMAT(r.hora, '%%H:%%i') AS hora,
               CONCAT_WS(' ', c.nombre, NULLIF(c.apellido, '')) AS cliente,
               c.telefono, b.nombre AS barbero,
               GROUP_CONCAT(s.nombre ORDER BY rs.id SEPARATOR ', ') AS servicios,
               CASE
                 WHEN r.estado IN ('No asistio', 'No asistió', 'No asistiÃ³', 'No asistiÃƒÂ³') THEN 'No asistio'
                 ELSE r.estado
               END AS estado,
               r.total
        FROM reservas r
        INNER JOIN clientes c ON c.id = r.id_cliente
        INNER JOIN barberos b ON b.id = r.id_barbero
        LEFT JOIN reserva_servicios rs ON rs.id_reserva = r.id
        LEFT JOIN servicios s ON s.id = rs.id_servicio
        {where}
        GROUP BY r.id, r.fecha, r.hora, c.nombre, c.apellido, c.telefono, b.nombre, r.estado, r.total
        ORDER BY r.fecha DESC, r.hora DESC
        """,
        tuple(params),
    )
    return cursor.fetchall()


@reports_bp.get("/export")
@require_roles(ROLE_ADMIN)
def export_report():
    """Exporta un reporte tabular compatible con Excel."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Fecha", "Hora", "Cliente", "Telefono", "Barbero", "Servicios", "Estado", "Total"])
    with db_cursor() as cursor:
        cursor.execute(
            """
            SELECT r.fecha, TIME_FORMAT(r.hora, '%H:%i') AS hora,
                   CONCAT_WS(' ', c.nombre, NULLIF(c.apellido, '')) AS cliente,
                   c.telefono, b.nombre AS barbero,
                   GROUP_CONCAT(s.nombre ORDER BY rs.id SEPARATOR ', ') AS servicios,
                   r.estado, r.total
            FROM reservas r
            INNER JOIN clientes c ON c.id = r.id_cliente
            INNER JOIN barberos b ON b.id = r.id_barbero
            LEFT JOIN reserva_servicios rs ON rs.id_reserva = r.id
            LEFT JOIN servicios s ON s.id = rs.id_servicio
            GROUP BY r.id, r.fecha, r.hora, c.nombre, c.apellido, c.telefono, b.nombre, r.estado, r.total
            ORDER BY r.fecha DESC, r.hora DESC
            """
        )
        for row in cursor.fetchall():
            writer.writerow(
                [
                    row["fecha"],
                    row["hora"],
                    row["cliente"],
                    row["telefono"],
                    row["barbero"],
                    row["servicios"],
                    row["estado"],
                    row["total"],
                ]
            )
    return Response(
        output.getvalue(),
        mimetype="application/vnd.ms-excel",
        headers={"Content-Disposition": "attachment; filename=reporte-barberia.xls"},
    )


@reports_bp.get("/export/xlsx")
@require_roles(ROLE_ADMIN)
def export_xlsx():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return fail("Falta instalar openpyxl. Ejecuta: pip install -r backend/requirements.txt", 503)

    where, params = filters()
    output = io.BytesIO()
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    def add_sheet(title, headers, rows):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in rows:
            ws.append(row)
        for column in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in column)
            ws.column_dimensions[column[0].column_letter].width = min(max(max_len + 2, 12), 42)

    with db_cursor() as cursor:
        reservations = fetch_reservation_rows(cursor, where, params)
        add_sheet(
            "citas",
            ["Fecha", "Hora", "Cliente", "Telefono", "Barbero", "Servicios", "Estado", "Total"],
            rows_as_lists(reservations, ["fecha", "hora", "cliente", "telefono", "barbero", "servicios", "estado", "total"]),
        )

        cursor.execute("SELECT id, nombre, precio, duracion_minutos, estado FROM servicios ORDER BY nombre")
        add_sheet("servicios", ["ID", "Nombre", "Precio", "Duracion", "Estado"], rows_as_lists(cursor.fetchall(), ["id", "nombre", "precio", "duracion_minutos", "estado"]))

        cursor.execute("SELECT id, nombre, telefono, estado FROM barberos ORDER BY nombre")
        add_sheet("barberos", ["ID", "Nombre", "Telefono", "Estado"], rows_as_lists(cursor.fetchall(), ["id", "nombre", "telefono", "estado"]))

        cursor.execute("SELECT id, nombre, cantidad, unidad, precio_venta, stock_minimo, estado FROM inventario ORDER BY nombre")
        add_sheet("inventario", ["ID", "Producto", "Stock", "Unidad", "Precio venta", "Stock minimo", "Estado"], rows_as_lists(cursor.fetchall(), ["id", "nombre", "cantidad", "unidad", "precio_venta", "stock_minimo", "estado"]))

        cursor.execute(
            """
            SELECT MAX(CONCAT_WS(' ', c.nombre, NULLIF(c.apellido, ''))) AS nombre, c.telefono,
                   COUNT(r.id) AS total_citas,
                   SUM(CASE WHEN LEFT(r.estado, 9) = 'No asisti' THEN 1 ELSE 0 END) AS strikes
            FROM clientes c
            INNER JOIN reservas r ON r.id_cliente = c.id
            GROUP BY c.telefono
            HAVING strikes > 0
            ORDER BY strikes DESC, total_citas DESC
            """
        )
        add_sheet("clientes con strikes", ["Nombre", "Telefono", "Citas", "Strikes"], rows_as_lists(cursor.fetchall(), ["nombre", "telefono", "total_citas", "strikes"]))

    del wb["Sheet"]
    wb.save(output)
    output.seek(0)
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=reporte-barberia.xlsx"},
    )


@reports_bp.get("/export/pdf")
@require_roles(ROLE_ADMIN)
def export_pdf():
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        where, params = filters()
        with db_cursor() as cursor:
            summary = summary_payload(cursor, where, params)
        return simple_pdf_response(
            [
                "Reporte Wichos Barber",
                f"Reservas: {summary['totalReservas']}",
                f"Confirmadas: {summary['reservasConfirmadas']}",
                f"Atendidas: {summary['reservasCompletadas']}",
                f"Canceladas: {summary['reservasCanceladas']}",
                f"No asistio: {summary['reservasNoShow']}",
                f"Ingresos citas: Q{summary['ingresosTotales']:.2f}",
                f"Ventas inventario: Q{summary['ingresosProductos']:.2f}",
            ]
        )

    where, params = filters()
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(output, pagesize=letter, title="Reporte Wichos Barber")
    story = [Paragraph("Reporte Wichos Barber", styles["Title"])]

    active_filters = {
        "fecha_inicio": request.args.get("fecha_inicio") or "sin filtro",
        "fecha_fin": request.args.get("fecha_fin") or "sin filtro",
        "barbero_id": request.args.get("barbero_id") or "todos",
        "estado": request.args.get("estado") or "todos",
        "servicio_id": request.args.get("servicio_id") or "todos",
    }
    story.append(Paragraph("Filtros usados: " + ", ".join(f"{k}: {v}" for k, v in active_filters.items()), styles["Normal"]))
    story.append(Spacer(1, 12))

    with db_cursor() as cursor:
        summary = summary_payload(cursor, where, params)
        metrics = [
            ["Reservas", summary["totalReservas"], "Confirmadas", summary["reservasConfirmadas"]],
            ["Atendidas", summary["reservasCompletadas"], "Canceladas", summary["reservasCanceladas"]],
            ["No asistio", summary["reservasNoShow"], "Ingresos citas", f"Q{summary['ingresosTotales']:.2f}"],
            ["Ventas inventario", f"Q{summary['ingresosProductos']:.2f}", "", ""],
        ]
        story.append(Table(metrics, colWidths=[110, 90, 110, 90]))
        story.append(Spacer(1, 14))

        def add_table(title, headers, rows):
            story.append(Paragraph(title, styles["Heading2"]))
            table_data = [headers] + rows[:12]
            table = Table(table_data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 12))

        add_table("Servicios mas solicitados", ["Servicio", "Total"], rows_as_lists(summary["porServicio"], ["servicio", "total"]))
        add_table("Desempeno por barbero", ["Barbero", "Citas", "Ingresos"], rows_as_lists(summary["porBarbero"], ["barbero", "total", "ingresos"]))
        add_table("Clientes frecuentes", ["Nombre", "Telefono", "Citas", "Strikes"], rows_as_lists(summary["clientesFrecuentes"], ["nombre", "telefono", "total_citas", "strikes"]))
        add_table("Stock bajo", ["ID", "Producto", "Stock", "Minimo"], rows_as_lists(summary["stockBajo"], ["id", "nombre", "cantidad", "stock_minimo"]))

    doc.build(story)
    output.seek(0)
    return Response(
        output.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": "attachment; filename=reporte-barberia.pdf"},
    )
